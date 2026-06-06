from flask import Blueprint, send_file, request, jsonify
from flask_login import login_required, current_user
from .models import db, Equipo, Movimiento
from io import BytesIO
from datetime import datetime
import json, os

bp = Blueprint("exports", __name__, url_prefix="/exportar")

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    EXCEL_OK = True
except ImportError:
    EXCEL_OK = False


@bp.route("/inventario")
@login_required
def inventario_excel():
    if not EXCEL_OK:
        return jsonify({"error": "openpyxl no disponible"}), 500

    estado  = request.args.get("estado", "")
    seccion = request.args.get("seccion", "")
    query   = Equipo.query
    if estado:  query = query.filter_by(estado=estado)
    if seccion: query = query.filter_by(seccion=seccion)
    equipos = query.order_by(Equipo.seccion, Equipo.icm).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario SISDAH"

    # Cabecera
    rojo   = "C8102E"
    blanco = "FFFFFF"
    gris   = "F5F5F5"
    cabecera = ["ICM","Marca","Modelo","Sección","Estado","S/N","Entrada","Retiro","Retirado por","Ticket","Pedido","Garantía","Ubicación","Notas"]

    for col, texto in enumerate(cabecera, 1):
        c = ws.cell(row=1, column=col, value=texto)
        c.font      = Font(bold=True, color=blanco)
        c.fill      = PatternFill("solid", fgColor=rojo)
        c.alignment = Alignment(horizontal="center")

    for row_i, eq in enumerate(equipos, 2):
        vals = [eq.icm, eq.marca, eq.modelo, eq.seccion, eq.estado, eq.sn,
                eq.fecha_entrada, eq.fecha_retiro, eq.retirado_por, eq.codigo_ticket,
                eq.numero_pedido, eq.garantia_fin, eq.ubicacion, eq.notas]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row_i, column=col, value=val or "")
            if row_i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=gris)

    # Anchos de columna
    anchos = [16,12,22,14,12,20,12,12,14,12,14,12,16,30]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = ancho

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fecha = datetime.now().strftime("%Y%m%d_%H%M")
    return send_file(buf, as_attachment=True,
                     download_name=f"inventario_SISDAH_{fecha}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/historial")
@login_required
def historial_excel():
    if not EXCEL_OK:
        return jsonify({"error": "openpyxl no disponible"}), 500

    movs = Movimiento.query.order_by(Movimiento.fecha.desc()).limit(1000).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historial"

    rojo  = "C8102E"
    blanc = "FFFFFF"
    cab   = ["Fecha","ICM","Tipo","Usuario DNI","Descripción","PC"]
    for col, texto in enumerate(cab, 1):
        c = ws.cell(row=1, column=col, value=texto)
        c.font = Font(bold=True, color=blanc)
        c.fill = PatternFill("solid", fgColor=rojo)

    for i, m in enumerate(movs, 2):
        ws.cell(i, 1, m.fecha.strftime("%d/%m/%Y %H:%M"))
        ws.cell(i, 2, m.icm)
        ws.cell(i, 3, m.tipo)
        ws.cell(i, 4, m.usuario_dni or "")
        ws.cell(i, 5, m.descripcion or "")
        ws.cell(i, 6, m.pc or "")

    buf = BytesIO(); wb.save(buf); buf.seek(0)
    fecha = datetime.now().strftime("%Y%m%d_%H%M")
    return send_file(buf, as_attachment=True,
                     download_name=f"historial_SISDAH_{fecha}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/importar_json", methods=["POST"])
@login_required
def importar_json():
    """Importa datos del stock.json de la app de escritorio."""
    if not current_user.es_admin:
        return jsonify({"ok": False, "error": "Sin permisos"})

    # Rutas posibles del JSON
    rutas = [
        r"C:\Users\rsacr\Desktop\PROYECTO HGUGM\app\data\stock.json",
        r"C:\Users\rsacr\Desktop\PROYECTO HGUGM\data\stock.json",
        r"D:\SISDAH_HGUGM_V.03\app\data\stock.json",
    ]
    ruta_json = None
    for r in rutas:
        if os.path.exists(r):
            ruta_json = r
            break

    if not ruta_json:
        return jsonify({"ok": False, "error": "No se encontró stock.json. Indica la ruta en Configuración."})

    try:
        with open(ruta_json, encoding="utf-8") as f:
            datos = json.load(f)
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)})

    insertados = duplicados = errores = 0
    for icm, eq in datos.items():
        if Equipo.query.filter_by(icm=icm).first():
            duplicados += 1
            continue
        try:
            nuevo = Equipo(
                icm=icm,
                marca=eq.get("marca"),
                modelo=eq.get("modelo"),
                seccion=eq.get("seccion"),
                estado=eq.get("estado", "en stock"),
                sn=eq.get("sn"),
                fecha_entrada=eq.get("fecha_entrada"),
                fecha_retiro=eq.get("fecha_retiro"),
                retirado_por=eq.get("retirado_por"),
                codigo_ticket=eq.get("codigo_ticket"),
                garantia_fin=eq.get("garantia_fin"),
                numero_pedido=eq.get("numero_pedido"),
                notas=eq.get("notas"),
            )
            db.session.add(nuevo)
            insertados += 1
        except Exception:
            errores += 1

    db.session.commit()
    return jsonify({"ok": True, "insertados": insertados,
                    "duplicados": duplicados, "errores": errores})
